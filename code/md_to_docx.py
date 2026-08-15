# -*- coding: utf-8 -*-
"""
Convert PAPER_DRAFT.md to a .docx manuscript.

Written because neither pandoc nor node is available on this machine; this is
a small converter for the constructs the draft actually uses (headings, pipe
tables, images, blockquotes, bullet/numbered lists, and inline bold / italic /
code), not a general Markdown implementation.

    python md_to_docx.py [input.md] [output.docx]

Images are resolved relative to the markdown file's directory.
"""

from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x60, 0x60, 0x60)


# --------------------------- inline formatting ---------------------------
TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")

# A line matching this starts a new block, so paragraph/list accumulation
# must stop before it.  The fenced-code marker belongs here: without it the
# paragraph accumulator swallows the fence and emits it as literal text.
BLOCK_START = r"^(#{1,4}\s|\||>|!\[|[-*]\s|\d+\.\s|-{3,}$|`{3})"


def add_runs(par, text, base_italic=False, color=None,
             bold=False, italic=False, code=False):
    """
    Emit runs for one line, honouring **bold**, *italic* and `code`.

    Recurses so that nested markers work — the figure table uses
    **`filename.pdf`**, and a single-pass parser strips the outer ** and then
    emits the backticks as literal characters.
    """
    parts = TOKEN.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_runs(par, part[2:-2], base_italic, color,
                     bold=True, italic=italic, code=code)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            add_runs(par, part[1:-1], base_italic, color,
                     bold=bold, italic=italic, code=True)
        elif (part.startswith("*") and part.endswith("*")
              and not part.startswith("**") and len(part) > 2):
            add_runs(par, part[1:-1], base_italic, color,
                     bold=bold, italic=True, code=code)
        else:
            r = par.add_run(part)
            r.bold = bold or None
            r.italic = True if (italic or base_italic) else None
            if code:
                r.font.name = MONO_FONT
                r.font.size = Pt(9)
            if color is not None:
                r.font.color.rgb = color


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


# ------------------------------- blocks ---------------------------------
def flush_table(doc, rows):
    """rows: list of list[str]; row 1 is the header, row 2 the --- separator."""
    rows = [r for r in rows if not re.fullmatch(r"[\s|:-]+", "|".join(r))]
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]

    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    total = Inches(6.4)
    col_w = int(total / ncol)
    t.autofit = False
    for row in t.rows:
        for c in row.cells:
            c.width = col_w

    for i, r in enumerate(rows):
        for j, txt in enumerate(r):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, txt.strip())
            for run in p.runs:
                run.font.size = Pt(8.5)
                # don't clobber a code run's monospace face
                if run.font.name != MONO_FONT:
                    run.font.name = BODY_FONT
                if i == 0:
                    run.bold = True
            if i == 0:
                shade(cell, "EEF2F7")
    doc.add_paragraph()


def hrule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "BBBBBB")
    bd.append(bottom)
    pPr.append(bd)


def convert(md_path, out_path):
    base = os.path.dirname(os.path.abspath(md_path))
    lines = open(md_path, encoding="utf-8").read().split("\n")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)      # US Letter
    for attr in ("left_margin", "right_margin"):
        setattr(sec, attr, Inches(1.0))
    sec.top_margin = sec.bottom_margin = Inches(0.9)

    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15

    tbl: list[list[str]] = []
    i = 0
    while i < len(lines):
        ln = lines[i]
        stripped = ln.strip()

        # ---- table accumulation ----
        if stripped.startswith("|"):
            tbl.append([c for c in stripped.strip("|").split("|")])
            i += 1
            continue
        if tbl:
            flush_table(doc, tbl)
            tbl = []

        if not stripped:
            i += 1
            continue

        # ---- fenced code block ----
        # Display formulas and code listings. Emitted verbatim in the mono
        # face, one paragraph per line, with inline markers left alone.
        if stripped.startswith("```"):
            i += 1
            body = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i].rstrip())
                i += 1
            i += 1                                   # closing fence
            while body and not body[0].strip():
                body.pop(0)
            while body and not body[-1].strip():
                body.pop()
            for k, ln in enumerate(body):
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Inches(0.28)
                pf.space_before = Pt(5 if k == 0 else 0)
                pf.space_after = Pt(5 if k == len(body) - 1 else 0)
                pf.line_spacing = 1.0
                r = p.add_run(ln if ln.strip() else " ")
                r.font.name = MONO_FONT
                r.font.size = Pt(8.8)
            continue

        # ---- horizontal rule ----
        if re.fullmatch(r"-{3,}", stripped):
            hrule(doc)
            i += 1
            continue

        # ---- image ----
        m = re.fullmatch(r"!\[(.*?)\]\((.+?)\)", stripped)
        if m:
            path = os.path.join(base, m.group(2))
            if os.path.exists(path):
                doc.add_picture(path, width=Inches(6.4))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p = doc.add_paragraph()
                add_runs(p, f"[missing image: {m.group(2)}]", color=MUTED)
            i += 1
            continue

        # ---- headings ----
        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(level=min(level, 4))
            h.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            h.paragraph_format.space_after = Pt(4)
            add_runs(h, m.group(2))
            for r in h.runs:
                r.font.name = BODY_FONT
                r.font.color.rgb = INK
                r.font.size = Pt({1: 17, 2: 14, 3: 11.5, 4: 10.5}[min(level, 4)])
            i += 1
            continue

        # ---- blockquote ----
        # Consecutive '>' lines are one quote. They have to be joined before
        # parsing, or an inline marker that spans the line break never closes
        # and the literal asterisks reach the document.
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, " ".join(buf).strip(), base_italic=True)
            continue

        # ---- lists ----
        # A list item may wrap over several source lines; its continuation
        # lines have to be joined before parsing, or an inline marker that
        # spans the break (**bold\ncontinued**) never closes and the literal
        # asterisks end up in the document.
        def gather(start):
            buf, j = [], start
            while j < len(lines):
                nxt = lines[j].strip()
                if j > start and (not nxt or re.match(
                        BLOCK_START, nxt)):
                    break
                buf.append(nxt)
                j += 1
            return " ".join(buf), j

        m = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if m:
            text, i = gather(i)
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s+", "", text))
            continue
        if re.match(r"^[-*]\s+", stripped):
            text, i = gather(i)
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^[-*]\s+", "", text))
            continue

        # ---- plain paragraph: join until a blank line ----
        buf = [stripped]
        j = i + 1
        while j < len(lines) and lines[j].strip() and not re.match(
                BLOCK_START, lines[j].strip()):
            buf.append(lines[j].strip())
            j += 1
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, " ".join(buf))
        i = j

    if tbl:
        flush_table(doc, tbl)

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "PAPER_DRAFT.md"
    dst = sys.argv[2] if len(sys.argv) > 2 else "PAPER_DRAFT.docx"
    print("wrote", convert(src, dst))
